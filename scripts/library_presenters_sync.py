#!/usr/bin/env python3
"""Refresh the presenter override table from Varya's sheet and report coverage.

The CMS has no presenter for the lecture programmes, so those episodes carry
no speaker at all and sit out the heaviest axis in the ranking. Varya filled
them in by hand («Заполнить ведущих», 86cbhq9q8). The sheet is shared with
named people rather than published, so there is no API key path to it: export
it (File, Download) and point this script at the download.

    scripts/library_presenters_sync.py ~/Downloads/lectures_with_presenter.xlsx
    scripts/library_presenters_sync.py --check

Only the `presenter` column is read. `кто_похож_на_лектора` beside it is a
generated guess, in the genitive and sometimes carrying a job title
(«Екатерины Михайловой», «Психотерапевт Алена»), and feeding that to a name
matcher would be worse than the gap it fills.

The table is keyed on content_id and lives beside the programme table rather
than being written back into `library_tags.jsonl`, which the tagging batch
regenerates: hand work a machine run can destroy is hand work that will be
lost.
"""

from __future__ import annotations

import argparse
import csv
import json
import sys
from pathlib import Path


REPO_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(REPO_ROOT / "src"))

DEFAULT_TABLE = REPO_ROOT / "data" / "library_presenters.csv"
DEFAULT_TAGS = REPO_ROOT / "data" / "library_tags.jsonl"
# The programme table decides whether a presenter counts as a speaker, so the
# coverage figure is wrong without it.
DEFAULT_PROGRAMS = REPO_ROOT / "data" / "library_programs.csv"

COLUMNS = ("content_id", "presenter")


def rows_from_export(path: Path) -> list[dict[str, str]]:
    """(content_id, presenter) pairs from an .xlsx or .csv export."""
    from rainrag.library_presenters import split_presenters

    if path.suffix.lower() in {".xlsx", ".xlsm"}:
        import openpyxl

        book = openpyxl.load_workbook(path, read_only=True, data_only=True)
        sheet = book[book.sheetnames[0]]
        raw = sheet.iter_rows(values_only=True)
        header = [str(c or "").strip() for c in next(raw, ())]
        records = [dict(zip(header, row, strict=False)) for row in raw]
    else:
        with open(path, encoding="utf-8-sig", newline="") as handle:
            reader = csv.DictReader(handle)
            if reader.fieldnames:
                reader.fieldnames = [(name or "").strip() for name in reader.fieldnames]
            records = list(reader)

    out: list[dict[str, str]] = []
    for record in records:
        # Excel hands back a float for a numeric cell, so 412707 arrives as
        # 412707.0 and would never match a content_id from the tag file.
        content_id = str(record.get("content_id") or "").strip()
        if content_id.endswith(".0"):
            content_id = content_id[:-2]
        names = split_presenters(record.get("presenter") or "")
        if content_id and names:
            out.append({"content_id": content_id, "presenter": ", ".join(names)})
    return out


def write_table(rows: list[dict[str, str]], path: Path) -> None:
    """Publish the table by rename, so a torn file is never read."""
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_suffix(path.suffix + ".tmp")
    with open(tmp, "w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=COLUMNS)
        writer.writeheader()
        writer.writerows(rows)
    tmp.replace(path)


def coverage(table: Path, tags: Path, programs: Path | None = None) -> tuple[int, int, int]:
    """(overrides, of those in the pool, of those that actually gain a speaker).

    The third number is the one that matters, and it is measured the way the
    Library measures it rather than guessed from the rows. Counting any row
    without CMS people overstated it twice: the tag file is append-only, so
    an older row still counted after a newer one gained a presenter, and an
    override on a programme whose genre demotes presenters was counted even
    though `resolve_speakers` leaves that episode with no speaker at all
    (CodeRabbit on #98). Resolving both ways over the deduplicated pool
    cannot drift from what the app does.
    """
    from rainrag.library_presenters import load_presenters
    from rainrag.library_programs import load_programmes
    from rainrag.library_similar import Episode, dedupe_latest

    overrides = load_presenters(table)
    if not tags.exists():
        return len(overrides), 0, 0
    programmes = load_programmes(programs) if programs else {}

    records = []
    for line in tags.read_text(encoding="utf-8").splitlines():
        if not line.strip():
            continue
        try:
            record = json.loads(line)
        except ValueError:
            continue
        if not record.get("error"):
            records.append(record)

    def pool(applied: dict[str, list[str]] | None) -> dict[str, Episode]:
        episodes = [Episode.from_record(r, programmes, applied) for r in records]
        return {e.video_hash: e for e in dedupe_latest(episodes)}

    without = pool(None)
    with_table = pool(overrides)
    in_pool = {h for h, e in with_table.items() if str(e.content_id or "") in overrides}
    gained = {
        h for h in in_pool if with_table[h].speakers and not without.get(h, with_table[h]).speakers
    }
    return len(overrides), len(in_pool), len(gained)


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("export", nargs="?", help="xlsx or csv export of the sheet")
    parser.add_argument("--table", default=str(DEFAULT_TABLE))
    parser.add_argument("--tags", default=str(DEFAULT_TAGS))
    parser.add_argument("--programs", default=str(DEFAULT_PROGRAMS))
    parser.add_argument("--check", action="store_true", help="report coverage only")
    args = parser.parse_args(argv)

    table = Path(args.table)
    if args.export and not args.check:
        rows = rows_from_export(Path(args.export))
        if not rows:
            print("no usable rows: expected columns content_id and presenter")
            return 1
        write_table(rows, table)
        print(f"wrote {len(rows)} presenter override(s) to {table}")
    elif not args.check:
        parser.error("give an export to sync, or --check to report coverage")

    total, in_pool, rescued = coverage(table, Path(args.tags), Path(args.programs))
    print(f"overrides: {total} | in the tagged pool: {in_pool} | gaining a speaker: {rescued}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
