#!/usr/bin/env python3
"""Refresh the programme table from Varya's sheet export and report coverage.

The Programs tab of the "Tags taxonomy" sheet is the editorial source of truth
for programme genre, and the speaker rule in `rainrag.library_programs` reads
it. The sheet is shared with named people rather than published, so there is no
API key path to it: export the tab as CSV (File, Download, Comma-separated
values) and point this script at the download.

    scripts/library_programs_sync.py ~/Downloads/Tags\\ taxonomy.xlsx

**Prefer the .xlsx export.** Varya marks the rows that are not really
programmes by colouring them, and a CSV export throws the colour away. Reading
the workbook recovers them, which matters for the coverage figure: counting
2,096 episodes of «(без программы)» and archive duplicates as uncovered
reported 83% for a table that actually reaches 98% of the material in scope.

A .csv export still works and leaves any existing exclusion list alone.

Usage:
    scripts/library_programs_sync.py <export.xlsx>   # table + exclusions
    scripts/library_programs_sync.py <export.csv>    # table only
    scripts/library_programs_sync.py --check         # report coverage only
"""

from __future__ import annotations

import argparse
import csv
import json
import shutil
import sys
from collections import Counter
from pathlib import Path


REPO_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(REPO_ROOT / "src"))

DEFAULT_TABLE = REPO_ROOT / "data" / "library_programs.csv"
DEFAULT_TAGS = REPO_ROOT / "data" / "library_tags.jsonl"
DEFAULT_EXCLUDED = REPO_ROOT / "data" / "library_programs_excluded.json"

# The bucket the catalogue uses for episodes with no programme at all. Varya
# colours it too, but it never appears in the `program` field by that name.
NO_PROGRAMME = "(без программы)"

REQUIRED_COLUMNS = {"title", "genre"}


def excluded_from_workbook(path: Path) -> list[str]:
    """Programme titles Varya has marked as not being programmes.

    She flags them by colouring the row, and every one of them also has an
    explanation in `comment`. The rule used here is "coloured and no genre",
    rather than matching one exact shade: a highlight she picks in a different
    tone still counts, while a row she has highlighted *and* given a genre to
    («ONLINE», at the time of writing) stays in scope.
    """
    try:
        import openpyxl
    except ImportError as exc:  # pragma: no cover - depends on the extra
        raise SystemExit(
            "Reading .xlsx needs the sheets extra: uv sync --extra sheets\n"
            "Or export the Programs tab as CSV, which skips the exclusions."
        ) from exc

    workbook = openpyxl.load_workbook(path)
    if "Programs" not in workbook.sheetnames:
        raise SystemExit(f"{path} has no 'Programs' sheet; tabs are {workbook.sheetnames}")
    sheet = workbook["Programs"]
    header = [str(c.value or "").strip() for c in sheet[1]]
    try:
        title_at, genre_at = header.index("title"), header.index("genre")
    except ValueError as exc:
        raise SystemExit(f"{path} Programs sheet is missing title or genre: {header}") from exc

    excluded: list[str] = []
    for row in sheet.iter_rows(min_row=2):
        title = row[title_at].value
        if not title:
            continue
        filled = row[0].fill.patternType == "solid"
        if filled and not (row[genre_at].value or "").strip():
            excluded.append(str(title).strip())
    return excluded


def _csv_from_workbook(path: Path, table: Path) -> None:
    """Write the Programs sheet out as the CSV the loader reads."""
    import openpyxl

    sheet = openpyxl.load_workbook(path)["Programs"]
    rows = list(sheet.iter_rows(values_only=True))
    header = [str(c or "").strip() for c in rows[0]]
    width = len([h for h in header if h])
    body = [r for r in rows[1:] if any(c not in (None, "") for c in r[:width])]
    with open(table, "w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(header[:width])
        for row in body:
            writer.writerow(["" if c is None else str(c) for c in row[:width]])
    print(f"wrote {table}: {len(body)} programmes")


def load_excluded(path: Path) -> set[str]:
    """Normalised titles that should not count towards coverage."""
    if not path.exists():
        return set()
    from rainrag.library_programs import normalise_title

    try:
        titles = json.loads(path.read_text(encoding="utf-8"))
    except ValueError:
        return set()
    return {normalise_title(str(t)) for t in titles if str(t).strip()}


def validate_export(path: Path) -> list[dict[str, str]]:
    """Read the export, failing loudly if it is not the Programs tab."""
    with open(path, encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        # Strip the header names here too, the way load_programmes does.
        # Validating against stripped names while the rows keep the padded
        # ones meant a header of " genre " passed the check and then read as
        # missing on every row, so the write report announced "0 with a
        # genre" for a table the loader reads perfectly well.
        if reader.fieldnames:
            reader.fieldnames = [(name or "").strip() for name in reader.fieldnames]
        rows = list(reader)
    if not rows:
        raise SystemExit(f"{path} has no data rows")
    missing = REQUIRED_COLUMNS - set(rows[0])
    if missing:
        raise SystemExit(
            f"{path} is missing the column(s) {sorted(missing)}. "
            "Export the Programs tab, not the whole workbook."
        )
    return rows


def latest_tag_rows(path: Path) -> list[dict[str, object]]:
    """Last good row wins, matching how the ranker reads the tagging output.

    Errors are dropped *before* the dedup, not after. The ranker does the same,
    and the difference is not cosmetic: an episode that was tagged, then
    re-tagged into a failure, still has a usable earlier row. Filtering last
    would discard it here while the interface keeps serving it, so the coverage
    report would blame the programme table for a gap that does not exist.
    """
    if not path.exists():
        return []
    latest: dict[str, dict[str, object]] = {}
    with open(path, encoding="utf-8") as handle:
        for line in handle:
            line = line.strip()
            if not line:
                continue
            try:
                record = json.loads(line)
            except ValueError:
                continue
            if not isinstance(record, dict) or record.get("error"):
                continue
            if record.get("video_hash"):
                latest[str(record["video_hash"])] = record
    return list(latest.values())


def report_coverage(table: Path, tags: Path, excluded_path: Path = DEFAULT_EXCLUDED) -> int:
    """Print how many tagged episodes the table reaches. Returns episodes missed.

    Coverage is measured against the material actually in scope. Rows Varya has
    marked as not programmes are taken out of the denominator rather than
    counted as gaps: with them in, the figure read 83% when the table reaches
    98% of the episodes it is supposed to.
    """
    from rainrag.library_programs import load_programmes, normalise_title

    programmes = load_programmes(table)
    excluded = load_excluded(excluded_path)
    episodes = latest_tag_rows(tags)
    if not episodes:
        print(f"no tagged episodes at {tags}; coverage not checked")
        return 0
    with_genre = 0
    out_of_scope = 0
    listed_no_genre: Counter[str] = Counter()
    unlisted: Counter[str] = Counter()
    for record in episodes:
        name = str(record.get("program") or "")
        key = normalise_title(name or NO_PROGRAMME)
        if key in excluded:
            out_of_scope += 1
            continue
        programme = programmes.get(key) if name else None
        if programme is None:
            unlisted[name or NO_PROGRAMME] += 1
        elif programme.genres:
            with_genre += 1
        else:
            listed_no_genre[programme.title] += 1
    total = len(episodes)
    in_scope = total - out_of_scope
    print(f"programmes in table: {len(programmes)}, marked not-a-programme: {len(excluded)}")
    print(f"tagged episodes: {total}")
    print(f"  out of scope (not a programme): {out_of_scope}")
    print(f"  in scope: {in_scope}")
    if in_scope:
        print(f"    has a genre: {with_genre} ({100 * with_genre / in_scope:.0f}%)")
    print(f"    listed without a genre: {sum(listed_no_genre.values())}")
    for name, count in listed_no_genre.most_common(5):
        print(f"        {count:>5}  {name}")
    print(f"    not in the table: {sum(unlisted.values())}")
    for name, count in unlisted.most_common(5):
        print(f"        {count:>5}  {name}")
    return sum(unlisted.values()) + sum(listed_no_genre.values())


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("export", nargs="?", help="CSV export of the Programs tab")
    parser.add_argument("--table", default=str(DEFAULT_TABLE))
    parser.add_argument("--tags", default=str(DEFAULT_TAGS))
    parser.add_argument("--excluded", default=str(DEFAULT_EXCLUDED))
    parser.add_argument(
        "--check", action="store_true", help="report coverage of the current table only"
    )
    args = parser.parse_args(argv)

    table = Path(args.table)
    excluded_path = Path(args.excluded)
    if not args.check:
        if not args.export:
            parser.error("give an export to load, or --check to only report coverage")
        from rainrag.library_programs import parse_genres

        source = Path(args.export)
        table.parent.mkdir(parents=True, exist_ok=True)
        if source.suffix.lower() == ".xlsx":
            excluded = excluded_from_workbook(source)
            excluded_path.write_text(
                json.dumps(excluded, ensure_ascii=False, indent=1), encoding="utf-8"
            )
            print(f"wrote {excluded_path}: {len(excluded)} rows marked not-a-programme")
            _csv_from_workbook(source, table)
        else:
            rows = validate_export(source)
            genres = sum(1 for row in rows if parse_genres(row.get("genre") or ""))
            shutil.copyfile(source, table)
            print(f"wrote {table}: {len(rows)} programmes, {genres} with a genre")
            print(f"(CSV export: cell colours are lost, {excluded_path.name} left as it was)")

    report_coverage(table, Path(args.tags), excluded_path)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
